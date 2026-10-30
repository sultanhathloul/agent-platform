"""Chunking and the XLSX header (owner decision A3): no chunk is longer than its budget, the
overlap included; a header longer than the budget is truncated and never squeezes out the
rows; title rows above an XLSX table are not its header. Arabic text throughout (G4)."""

from app.context.tokens import default_counter
from app.rag.chunk import chunk_segments
from app.rag.config import load_docs_config
from app.rag.extract import extract
from app.rag.types import Segment
from tests.fixtures.documents import labour_rows_xlsx_bytes, xlsx_bytes

import pytest


@pytest.fixture(scope="module")
def docs_config(config_dir):
    return load_docs_config(config_dir)


class _Words:
    """One token per word: chunk sizes a test can state exactly."""

    def count(self, text: str) -> int:
        return len(text.split())


def test_the_overlap_gives_way_so_no_chunk_passes_the_budget():
    """Owner decision A3: after a cut, the overlap (25) plus the next unit (190) would make 215."""
    units = [" ".join(["أ"] * 100), " ".join(["ب"] * 25), " ".join(["ج"] * 190)]
    chunks = chunk_segments(
        [Segment(text="\n".join(units), page=1)], _Words(), max_tokens=200, overlap_tokens=30
    )
    assert [c.tokens for c in chunks] == [125, 190]
    assert all(c.tokens <= 200 for c in chunks)


def test_a_header_longer_than_the_budget_is_truncated_and_rows_still_fit():
    header = " ".join(f"عمود{i}" for i in range(300))
    rows = "\n".join(f"صف {i} قيمة" for i in range(60))
    chunks = chunk_segments(
        [Segment(text=rows, section="ورقة", repeat=header)], _Words(), max_tokens=200, overlap_tokens=30
    )
    assert all(c.tokens <= 200 for c in chunks)
    assert all(c.text.startswith("عمود0 عمود1") for c in chunks)
    assert len(chunks) < 10  # each chunk holds rows: the budget did not collapse to one token
    assert sum(c.text.count("صف ") for c in chunks) >= 60


def test_real_token_counts_stay_within_the_budget_on_arabic_text():
    counter = default_counter()
    paragraphs = [
        ("يلتزم الموظف بالحضور في الوقت المحدد ويحافظ على سرية المعلومات " * n).strip()
        for n in (1, 9, 2, 14, 1, 13)
    ]
    chunks = chunk_segments(
        [Segment(text="\n\n".join(paragraphs), page=1)], counter, max_tokens=200, overlap_tokens=30
    )
    assert all(c.tokens <= 200 for c in chunks)


def test_xlsx_chunks_repeat_the_sheet_the_tables_titles_and_its_header(docs_config):
    """ADR-0410 (superseding A3's header-only repeat): every chunk opens with the sheet's name and every row
    above the first data row, titles and header together; the rows below are the text."""
    rows = [
        ["ميزانية هيئة الواحة لعام 2031"],
        ["أُعدت في 1 محرم"],
        ["البند", "المبلغ", "الملاحظة"],
        ["التدريب", 120000, "دورات داخلية"],
        ["السفر", 45000, "مهام رسمية"],
    ]
    [segment] = extract("xlsx", xlsx_bytes({"الميزانية": rows}), docs_config).segments
    assert segment.repeat == ("[الميزانية]\nميزانية هيئة الواحة لعام 2031\nأُعدت في 1 محرم\nالبند | المبلغ | الملاحظة\n"
                              "[columns] البند | المبلغ | الملاحظة")
    assert segment.text.split("\n") == ["التدريب | 120000 | دورات داخلية", "السفر | 45000 | مهام رسمية"]


def test_a_statistical_tables_multi_row_header_and_title_open_every_chunk(docs_config):
    """Table 3-1 of the Q2 2026 labour-market workbook, shortened: a title, the table number, three header
    rows (nationality, then sex) and a header of years left as a header."""
    rows = [
        ["", "إحصاءات سوق العمل الربع الثاني 2026"],
        ["", "التوزيع النسبي للمشتغلين بحسب الجنسية و الجنس(%) - سلسلة زمنية"],
        ["جدول (3-1)"],
        ["الربع / السنة", "الجنسية"],
        ["", "السعوديين", "", "غير السعوديين"],
        ["", "الذكور", "الاناث", "الذكور", "الاناث"],
        *[[f"الربع {q} 2021", 70.4, 29.5, 90.4, 9.5] for q in range(1, 40)],
        ["الربع الثاني 2026", 66.366373, 33.633627, 91.46388445, 8.536115554],
    ]
    [segment] = extract("xlsx", xlsx_bytes({"3-1": rows}), docs_config).segments
    assert segment.repeat.split("\n")[:3] == ["[3-1]", " | إحصاءات سوق العمل الربع الثاني 2026",
                                               " | التوزيع النسبي للمشتغلين بحسب الجنسية و الجنس(%) - سلسلة زمنية"]
    assert " | الذكور | الاناث | الذكور | الاناث\n[columns] " in segment.repeat
    chunks = chunk_segments([segment], _Words(), max_tokens=120, overlap_tokens=10)
    last = next(c for c in chunks if "الربع الثاني 2026 | 66.366373" in c.text)
    assert "التوزيع النسبي للمشتغلين" in last.text and "الذكور | الاناث" in last.text
    years = [["جدول (2)"], ["المؤشر", "2025", "2026"], ["معدل البطالة", 3.2, 3.0]]
    [segment] = extract("xlsx", xlsx_bytes({"2": years}), docs_config).segments
    assert segment.repeat.endswith("المؤشر | 2025 | 2026\n[columns] المؤشر | 2025 | 2026")
    assert segment.text == "معدل البطالة | 3.2 | 3"


def test_a_sheet_without_a_table_takes_its_first_row(docs_config):
    [segment] = extract(
        "xlsx", xlsx_bytes({"ملاحظات": [["ملاحظة أولى"], ["ملاحظة ثانية"]]}), docs_config
    ).segments
    assert segment.repeat == "[ملاحظات]\nملاحظة أولى" and segment.text == "ملاحظة ثانية"


def _gastat_table() -> bytes:
    """Table 1 of the Q2 2026 labour-market workbook, as GASTAT builds it: a title merged over the table, the
    nationality merged over its three columns, the cells stored at full precision with a one-decimal format."""
    import io

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "1"
    sheet["B1"] = "المؤشرات الرئيسية لاحصاءات سوق العمل حسب الجنسية والجنس"
    sheet.merge_cells("B1:G1")
    sheet["A2"], sheet["B2"], sheet["E2"] = "المؤشرات", "السعوديين", "غير السعوديين"
    sheet.merge_cells("B2:D2")
    sheet.merge_cells("E2:G2")
    for col, sex in zip("BCDEFG", ["الذكور", "الاناث", "الإجمالي"] * 2, strict=True):
        sheet[f"{col}3"] = sex
    values = [4.84352836909752, 9.56074240840159, 6.48407237148921, 1.09065752454951, 3.81932584269662,
              1.32960919200363]
    sheet["A4"] = "معدل البطالة"
    for col, value in zip("BCDEFG", values, strict=True):
        sheet[f"{col}4"] = value
        sheet[f"{col}4"].number_format = "[$-10401]#,##0.0;\\(#,##0\\)"
    sheet["A5"], sheet["B5"], sheet["C5"] = "عدد المشتغلين", 12345678, 0.256
    sheet["B5"].number_format, sheet["C5"].number_format = "#,##0", "0.0%"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_merged_headers_name_every_column_and_cells_read_as_displayed(docs_config):
    """ADR-0413: «غير السعوديين» merged over three columns names each of them, each column's full label is
    written once, and a cell reads as the workbook shows it, not at full float precision."""
    [segment] = extract("xlsx", _gastat_table(), docs_config).segments
    lines = segment.repeat.split("\n")
    assert lines[1] == " | المؤشرات الرئيسية لاحصاءات سوق العمل حسب الجنسية والجنس"  # a title is not repeated
    assert "المؤشرات | السعوديين | السعوديين | السعوديين | غير السعوديين | غير السعوديين | غير السعوديين" in lines
    assert lines[-1] == ("[columns] المؤشرات | السعوديين › الذكور | السعوديين › الاناث | السعوديين › الإجمالي | "
                         "غير السعوديين › الذكور | غير السعوديين › الاناث | غير السعوديين › الإجمالي")
    assert segment.text.split("\n") == ["معدل البطالة | 4.8 | 9.6 | 6.5 | 1.1 | 3.8 | 1.3",
                                        "عدد المشتغلين | 12,345,678 | 25.6%"]


def test_number_formats_as_a_spreadsheet_shows_them():
    from app.rag.extract.xlsx import displayed

    assert displayed(83.60299013, "[$-10401]#,##0.0;\\(#,##0\\)") == "83.6"
    assert displayed(1234.5, "#,##0") == "1,235" and displayed(0.4567, "0.00%") == "45.67%"
    assert displayed(12.5, '0.0" %"') == "12.5 %" and displayed(3.0, "General") == "3"
    assert displayed(83.60299013456789, "General") == "83.60299013" and displayed(46000, "dd/mm/yyyy") == "46000"


def test_a_title_merged_over_part_of_the_table_is_written_once(docs_config):
    """Sheet 2-3 of the Q2 2026 workbook: its title is merged over fewer columns than the table holds; it is
    a title, written once, not once per column (ADR-0413, extractor version 4)."""
    import io

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "2-3"
    sheet["B1"] = "معدل البطالة حسب الجنسية والجنس والمستوى التعليمي"
    sheet.merge_cells("B1:D1")
    sheet["A2"], sheet["B2"], sheet["C2"], sheet["D2"], sheet["E2"] = "التعليم", "الذكور", "الاناث", "الإجمالي", "ملاحظة"
    sheet.append(["ثانوي", 3.1, 9.4, 4.2, "تقديري"])
    buffer = io.BytesIO()
    book.save(buffer)
    [segment] = extract("xlsx", buffer.getvalue(), docs_config).segments
    assert segment.repeat.split("\n")[1] == " | معدل البطالة حسب الجنسية والجنس والمستوى التعليمي"



def test_a_large_workbook_streams_within_the_sandbox_memory(docs_config):
    """Review round 1 of PR #17, finding 8 (ADR-0422): a 120k-row workbook (6.3 MB) peaked at 370 MB of
    Python heap opened in full, past what the sandbox's 640 MiB of address space leaves; streamed
    read-only with its merges read from the sheet XML it stays near the sheet's text."""
    import tracemalloc

    from app.rag.extract.xlsx import XlsxExtractor

    data = labour_rows_xlsx_bytes(120_000)
    tracemalloc.start()
    try:
        document = XlsxExtractor().extract(data, docs_config)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    print(f"xlsx {len(data) / 1e6:.1f} MB, 120000 rows: peak heap {peak / 2**20:.0f} MiB")
    assert peak < 160 * 2**20, peak
    [segment] = document.segments
    assert "[columns] المؤشر | السعوديين › الذكور" in segment.repeat
    assert segment.text.count("\n") == 120_000 - 1


def test_a_merge_declared_over_a_million_rows_fills_only_the_rows_the_sheet_holds(docs_config):
    """A merged range costs the rows that exist, never its declared span (ADR-0422). The ranges are written
    into the sheet XML directly: openpyxl itself would build a cell per merged position."""
    import io
    import time
    import zipfile

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["المنطقة", "القيمة", "النسبة"])
    sheet.append(["الرياض", 10, 2.5])
    sheet.append([None, 11, 2.6])
    plain = io.BytesIO()
    book.save(plain)
    merges = '<mergeCells count="2"><mergeCell ref="A2:A1000000"/><mergeCell ref="D1:XFD1048576"/></mergeCells>'
    out = io.BytesIO()
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = body.replace(b"</sheetData>", b"</sheetData>" + merges.encode(), 1)
            target.writestr(item, body)
    started = time.monotonic()
    [segment] = extract("xlsx", out.getvalue(), docs_config).segments
    assert time.monotonic() - started < 5
    assert segment.text.split("\n") == ["الرياض | 10 | 2.5", "الرياض | 11 | 2.6"]


def test_the_xlsx_extractor_version_moves_so_the_seed_reindexes():
    from app.rag.extract import extractor_version

    assert extractor_version("xlsx") == "5"
